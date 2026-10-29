"""Build a source-grounded, text-only project reference without publishing secrets."""
from __future__ import annotations

import ast
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public' / 'about_system'
OUT.mkdir(parents=True, exist_ok=True)
sections: list[tuple[int, str, list[str]]] = []


def section(title, text, level=1):
    paragraphs = [re.sub(r'\s+', ' ', p).strip() for p in text.strip().split('\n\n') if p.strip()]
    sections.append((level, title, paragraphs))


section('1. Document scope, evidence, and reading guide', '''
Intrusion Detection and Monitoring System: comprehensive implementation and operations reference. Prepared on 29 October 2026, using Asia/Manila as the document date context. This document describes the local project snapshot in the Intrusion-Detection-Monitoring-System repository. It is written entirely as searchable text, without screenshots, diagrams, or source-code listings. It includes detailed implementation explanations, an endpoint register, a database dictionary, operational procedures, a source member register, research material extracts, and a file inventory.

The primary evidence is executable application source: Python services under setup, Laravel controllers, middleware, models, routes, migrations, views, configuration, tests, dependency manifests, and the API ZIP archive. Training manuscripts and notebooks are supplementary historical evidence. When a manuscript, comment, or older artifact differs from active source, this reference identifies the difference rather than combining both as if they were one implementation. The archived Render API is a separate historical path, not proof of the deployment currently running on Render.

The local source can explain intended behavior, defaults, and implementation limitations. It cannot independently establish which revision is deployed on Hostinger, whether Python services are currently running, the contents of the production database, current router firewall state, actual production accuracy, or subscriber connectivity. The user's hosting screenshots establish that the role migration completed; they do not verify every deployed component. No live router operations, production schema changes, traffic generation, model retraining, or production database queries are performed to produce this document.

Secrets are deliberately excluded. Environment variable names are documented, but values from private environment files, passwords, application encryption keys, authentication tokens, connection URLs containing credentials, and runtime session contents are not included. Binary models are inventoried without deserializing them. Dependency and generated environment internals are described by category rather than reproduced as application source. Source filenames, method names, schema names, and endpoint paths remain explicit because they are necessary for maintenance.

Read sections two through seven for architecture and runtime flow; sections eight through thirteen for website, API, database, and training; sections fourteen through eighteen for deployment, testing, limitations, and maintenance. The remaining registers provide traceability to individual source members, migration definitions, notebook contents, historical research documents, and project files. The accompanying plain-text document contains the same narrative and is provided for editing and searching.
''')
section('2. Purpose, actors, and component boundaries', '''
The project combines network intrusion detection with human-directed subscriber security management. It collects MikroTik NetFlow export records, associates observed addresses with PPPoE subscriber names, computes numerical traffic features, sends those features to a two-stage classifier, stores classifications in MySQL, and presents the results through a Laravel website. An administrator can request blocking or restoration; a local router-facing reconciliation agent implements those requests. Detection alone does not automatically block a subscriber in the current main path.

The monitored subscriber generates network traffic through the MikroTik gateway. The gateway exports NetFlow records and exposes active PPP sessions and firewall management through the RouterOS API. The local monitoring computer receives those records, performs feature aggregation, runs or contacts the detector, and reaches the database and router. Hostinger hosts the Laravel website and the shared MySQL database. A browser displays database-derived dashboards and submits administrator actions. The database is the coordination point between detection, presentation, and enforcement.

The administrator has read access plus threat management permissions. The viewer account has read access to protected pages and read endpoints, including report export, but cannot submit threat changes. The detector is a machine client of the model service and a database writer. The blocker/restorer is a database reader and router controller. These machine processes do not authenticate as the website viewer; website session restrictions do not automatically secure their separate network interfaces.

The principal local scripts are setup/Mikrotek PPPoe Sniffer.py, setup/Detection Process.py, and setup/Pending Blocker.py. The historical spelling Mikrotek is part of the filename. The first script is a NetFlow collector rather than a raw packet sniffer. The second hosts a Flask inference API and writes classifications directly to MySQL. The third performs both blocking and restoration. No separate current Restorer.py service is required or present in the inspected main setup.

Laravel follows a route, middleware, controller, model, and Blade view structure. Routes select controllers; CheckSession validates the custom session and role; controllers query or update Eloquent models; Blade templates render HTML; browser JavaScript polls selected endpoints. Database migrations define persistent storage. Composer manages PHP dependencies; npm and Vite manage the source asset build; most visible administration pages also load prebuilt static template assets directly from public/admin.
''')
section('3. End-to-end system flow and timing', '''
Step one: configure MikroTik to export NetFlow version five records to the monitoring computer on UDP port 2055 and make RouterOS API access available to that computer. The repository contains router inspection and removal notes but does not automatically configure Traffic Flow export. The collector binds an IPv4 UDP socket on all local interfaces and accepts packets only when the sender address equals the configured MikroTik host.

Step two: query active PPP sessions through the RouterOS API. The collector builds a mapping from each current subscriber address to its stable PPPoE username, filters excluded address prefixes, and adds recognized addresses to a round-robin queue. A current address is an observation, not the long-term subscriber identity. The queue is refreshed using active session data, and addresses absent from a successful refreshed mapping are removed from queue state.

Step three: parse each supported NetFlow packet. Version five uses a 24-byte header followed by 48-byte records. Extract source and destination IPv4 addresses, source and destination ports, packet count, byte count, first and last flow timestamps, TCP flag bits, and protocol number. Record a local collection timestamp. Store each accepted flow in both source-address and destination-address indexes.

Step four: approximately every 60 seconds, select one queued subscriber and gather records where that subscriber appears in either direction. Require at least ten indexed flow records before analysis. Retain at most the most recent 200 outgoing records and 200 incoming records for feature calculation. Refresh subscriber mappings approximately every 300 seconds and remove stored flow records older than one hour during the periodic analysis branch.

The timing checks are inside the successful UDP receive path. A socket timeout immediately continues the loop. Consequently, these are traffic-driven checks, not independent scheduled timers: no incoming export traffic can delay subscriber refresh, analysis, and memory cleanup. A queue with N active clients normally revisits one client in roughly N times 60 seconds, plus processing and traffic delays; it does not analyze every subscriber every 60 seconds.

Step five: submit one JSON request to the configured inference URL, with a 30-second request timeout. The current detector validates presence of all required model features, coerces feature values to finite numeric values, computes Stage A attack probability, applies the configured threshold, and invokes Stage B only when Stage A predicts attack. The response includes the decision, probability, optional attack family, used threshold, and feature accounting information.

Step six: append the classification record to a local JSON-lines prediction log and attempt a database transaction. A benign result creates a detection row without an incident. An attack result either updates an eligible incident or creates a new unresolved incident, then inserts the detection referencing that incident. The website reads the same database and updates selected displays through polling or page navigation.

Step seven: an administrator selects Block. The website updates the chosen incident to blocked, records the request timestamp, and clears restoration and synchronization markers. The browser response confirms the database request, not router completion. Approximately every 20 seconds, the local agent reads all blocked PPPoE subscribers, resolves each current address, and maintains subscriber-owned forward-chain drop rules.

Step eight: an administrator selects Restore. The controller method is historically named ignore, but the operation sets the incident to resolved and records restoration time. The local agent only processes a pending restore when there is no other blocked incident for the same subscriber. It removes subscriber-owned rules, clears tracked connections for relevant addresses, and marks the resolved incident synchronized. The incident and its logs remain as historical evidence.

The website dashboard polls every 15 seconds while the document is visible. The system-status alert fragment also polls every 15 seconds. The live-monitor demonstration updates synthetic chart values every two seconds and timestamps every five seconds; those intervals do not measure the real collector or router. Overall response time combines export timing, queue position, inference time, database availability, browser polling, and enforcement polling.
''')
section('4. NetFlow collector and subscriber discovery', '''
Source: setup/Mikrotek PPPoe Sniffer.py. MikroTikConnector holds a RouterOS connection pool, API handle, and connection state. connect creates a pool using the configured router host, API port, username, and password. The current code selects no SSL and plaintext login. A failed connection is logged and reported as false. get_pppoe_clients reconnects when necessary, reads /ppp/active, extracts address and name, discards incomplete entries, and returns the accepted address-to-name mapping.

The is_client_ip helper rejects empty addresses, zero and broadcast endpoints, and a literal list of infrastructure-like string prefixes. The list includes the 100.64 through 100.127 carrier-grade address prefixes, selected 190.168 through 193.168 prefixes, loopback, link-local, and several multicast or reserved prefixes. It uses textual starts-with matching rather than a formal subnet library. It should not be described as a complete standards-based IPv4 classification policy. The separate is_infrastructure_ip helper also excludes addresses beginning with 192.168 and selected multicast or reserved prefixes.

NetFlowCollector creates the UDP socket, sets a one-second receive timeout, initializes per-source and per-destination flow lists, initializes the round-robin queue and known-client set, and maintains the current PPPoE mapping. parse_netflow_v5 checks minimum packet length and stops when an advertised record would exceed the available bytes. The main loop checks the NetFlow version before calling it. Versions other than five are ignored; template-based NetFlow version nine and IPFIX are not implemented by this parser.

Flow duration is the exported last timestamp minus the first timestamp in milliseconds. Negative or wrapped exporter counters are not explicitly repaired. The parser excludes a record if either endpoint is considered infrastructure by its filter. This may exclude traffic that an operator expects to monitor, particularly private-address traffic through a gateway. Every surviving record receives a local ISO-format timestamp; the collector does not store packet payloads or reconstruct application content.

update_client_list replaces the address-to-name mapping, prunes addresses absent from the new active set, and queues newly recognized addresses. get_next_client rotates the deque before returning its first element. store_flow indexes one flow under both endpoints, and get_client_flows concatenates both indexes for the selected address. The indexes intentionally support bidirectional observation, but they are not a strict five-tuple conversation reconstruction. Self-addressed records can appear in both indexes.

In the receive loop, subscriber discovery is refreshed only if the router query returns a nonempty mapping. An empty mapping or failed query therefore leaves the prior mapping in place. This is a relevant limitation when every PPP session disconnects or API access fails. The receive buffer is 8192 bytes. Export datagrams larger than that buffer can be truncated, while the parser processes only records that fit in the received buffer.

analyze_next_client consumes the next queue position even when it skips a subscriber with too few records. It builds a payload, logs traffic summaries, and posts it once. There is no persistent outbound queue, request retry loop, duplicate identifier, or delivery acknowledgement tied to a committed database row. Data retained in memory may be re-aggregated on a later turn, so successive classifications are overlapping observations, not necessarily disjoint network sessions.

The collector writes setup/logs/merged_ids.log and console output. MergedIDS obtains an initial PPP client mapping before entering the receive loop. The Windows launcher opens each component in a separate console. Exceptions are printed with a pause-on-error hook where applicable. Stopping the console or losing the monitoring computer stops collection; the hosted website can remain available while its data becomes stale.
''')
section('5. Traffic features and measurement semantics', '''
The current training contract contains fourteen ordered features under schema version mikrotik-netflow-v1. Feature order is taken from the loaded model bundle during inference, rather than inferred from JSON key order. Source and destination addresses, subscriber identity, ports, protocol, timestamps, TCP-flag counters, and several extra statistics are also transmitted for context, but not all are model inputs or persisted database columns.

destination_port is the latest context record's destination port. flow_duration is the sum of selected flow durations in milliseconds multiplied by one thousand, producing a microsecond representation. It is a sum of exporter durations across selected records, not the wall-clock span between first and last client activity. Concurrent flows can therefore contribute overlapping time.

total_fwd_packets and total_backward_packets are packet totals for outgoing and incoming selected records. total_length_of_fwd_packets and total_length_of_bwd_packets are their byte totals. fwd_packet_length_mean and bwd_packet_length_mean are packet-weighted means derived from each export record's bytes divided by packets. These represent aggregate estimates; NetFlow records do not supply each original packet length.

flow_bytes_s and flow_packets_s divide aggregate bytes and packets by the summed duration converted to seconds. fwd_packets_s and bwd_packets_s use the same denominator for each directional packet count. If aggregate duration is not positive, the code uses one second as a fallback denominator. down_up_ratio is incoming packets divided by the greater of outgoing packets and one. average_packet_size is aggregate bytes divided by the greater of aggregate packets and one.

The collector includes both compact database-friendly keys and normalized training keys. tot_fwd_pkts corresponds to total_fwd_packets; tot_bwd_pkts corresponds to total_backward_packets; tot_fwd_bytes corresponds to total_length_of_fwd_packets; tot_bwd_bytes corresponds to total_length_of_bwd_packets; fwd_pkt_len_mean and bwd_pkt_len_mean correspond to their expanded packet-length feature names; flow_pkts_s corresponds to flow_packets_s; dst_port corresponds to destination_port.

Additional transmitted fwd_iat_mean and bwd_iat_mean are computed by dividing the summed duration in milliseconds by the corresponding directional packet count with a minimum denominator of one. They are not measured packet-by-packet inter-arrival intervals, and they are not part of the fourteen-feature current training contract. syn_count, ack_count, fin_count, and rst_count count selected export records carrying each flag bit, rather than reconstructing the exact number of packets with that flag.

The context record prefers the latest outgoing record and falls back to the latest incoming one. src_ip is always set to the subscriber being analyzed, even when the fallback context record originally had that subscriber as its destination. Thus destination address and port metadata may not describe a canonical subscriber-to-server conversation in the incoming-only case. A classification represents a subscriber aggregate across many records and destinations, not proof that one displayed destination caused the prediction.

Packet-length means are calculated by expanding repeated mean values into lists once per packet. Large exported packet counts can consume substantial memory even though only a few hundred flow records are selected. A weighted-sum calculation would avoid this implementation cost, but this reference records the present behavior without changing it. Training data uses CICFlowMeter-derived columns, while deployment computes aggregate NetFlow approximations; matching column names and units does not alone prove identical statistical distributions.
''')
section('6. Detector API, persistence, and incident assignment', '''
Source: setup/Detection Process.py. The script loads environment settings from setup/.env, chooses a model filename under setup/data/model, and deserializes a trusted joblib bundle at startup. The default filename is netflow_ids.joblib. Required bundle entries are bin for the binary model, mul for attack-family classification, features for ordered feature names, and best_threshold. An optional schema_version is exposed to clients; missing schema metadata becomes unknown.

IDS_THRESHOLD overrides the bundle threshold and must be within zero and one inclusive. The code trims expected feature names and incoming JSON keys. It maps selected compact sender names into normalized model names when the normalized field is absent. A supplied normalized field takes precedence over its alias. Missing required feature names produce a 422 response containing error missing_model_features, schema_version, missing_expected_features, present_expected_features, and expected_features_count.

For each required feature, conversion to float is attempted. Conversion failures, NaN, and infinity become zero. Therefore feature presence validation is stricter than value validation: a present but malformed value may silently become zero. The API expects an object-like JSON body; request JSON is forced, and there is no dedicated input schema enforcing valid ranges, record size, subscriber identity, or approved source hosts. The source does not configure API authentication or rate limiting.

Stage A calls predict_proba and takes the second probability column as the attack probability. A probability at or above the selected threshold yields risk_level attack; lower values yield benign. Stage B predicts one attack-family string only for an attack. Benign attack_type remains null. The response returns risk_level, prob_attack, attack_type, used_threshold, present_expected_features, missing_expected_features, received_features_count, api_version debug-v3, and schema_version.

The detector constructs a record containing an epoch-second timestamp, the received normalized data, and prediction results. append_log creates the configured log directory and appends a JSON record to predictions.jsonl. Log failures are printed and do not stop database saving. save_to_database constructs a selected column payload, obtains a MySQL connection with autocommit disabled, assigns an incident if appropriate, inserts the detection, and commits. On transactional failure it rolls back and prints an error; the outer inference handler can still return a successful classification response.

A successful inference response is consequently not a durable-storage guarantee. Missing database name or username explicitly skips insertion. Database exceptions are handled internally rather than converted into an HTTP failure. Neither a local prediction-log entry nor a successful collector response proves that Laravel can read a newly committed row. Operators should verify database timestamps or the logs page independently when diagnosing missing observations.

assign_incident returns no incident for non-attack records. For an attack with a PPPoE name, it searches for the same subscriber name and attack type. A blocked incident is eligible without the unresolved time-window restriction and is preferred over unresolved incidents. Otherwise an unresolved incident is eligible only if its last detection falls within INCIDENT_WINDOW_MINUTES, which defaults to fifteen and is clamped to at least one. The selected row is locked for update within the transaction.

An eligible incident increments event_count and updates last_detected_at and updated_at. Otherwise a new incident starts with status unresolved, event_count one, and first/last detection timestamps. Resolved incidents are never reopened by this query. An attack without a PPPoE name creates a new incident rather than reusing an unidentified address group. The new detection is linked through incident_id. The schema does not enforce a uniqueness constraint covering eligible incident identity; row locking does not fully prevent concurrent first-insert races when no matching incident exists.

The script starts Flask on all interfaces, port 5000, with debug enabled when executed directly. This is an implementation fact, not a hardened production deployment recommendation. The API root returns a simple running message; it does not exercise the model prediction path, database transaction, collector, or router. Loading an artifact can fail before the server starts. The bundle is not inspected dynamically by this documentation process because joblib deserialization executes trusted Python object loading behavior.
''')
section('7. Blocker, restorer, and router reconciliation', '''
Source: setup/Pending Blocker.py. The process loads shared settings, creates setup/logs/blocked_threat_sync.log, connects to RouterOS, and repeatedly polls MySQL. BLOCK_SCAN_INTERVAL defaults to twenty seconds. BLOCK_RULE_PREFIX defaults to IDS block. DB_TABLE defaults to intrusion_logs and DB_INCIDENT_TABLE defaults to threat_incidents. This service must run on a host that can reach both the database and the MikroTik API.

get_blocked_subscribers selects distinct nonempty PPPoE names with at least one blocked incident. It intentionally continues enforcing a blocked subscriber on every scan rather than treating a block as a one-time command. get_pending_restores selects resolved incidents whose firewall_synced_at is null, whose PPPoE name is nonempty, and for which no blocked incident exists for that same name. This prevents restoring one historical incident from overriding another currently blocked incident for the subscriber.

The main loop processes all blocked subscribers before pending restores, isolates individual subscriber failures with logging, catches scan-level errors, and sleeps before the next iteration. Initial ensure_connected occurs before the main loop; a startup router connection failure can terminate startup. Inside the loop, ensure_connected primarily checks whether an API object exists rather than proving an existing session is healthy. There is no general reconnect/backoff supervisor implemented around every transport failure.

Stable rule ownership is encoded as a comment consisting of the configured prefix followed by pppoe and the subscriber name, with a colon between pppoe and the name. The comment is truncated to 255 characters. Matching rules are sought in the forward chain with action drop and this exact ownership comment. These are source-address drop rules, not PPP account deletion, password changes, or a complete bidirectional quarantine policy. Existing firewall rule order and FastTrack behavior remain relevant to effective enforcement.

resolve_current_ip_for_pppoe queries /ppp/active using the supplied name. If no exact response is returned, it reads active sessions and compares trimmed lowercased names. Missing active sessions or missing assigned addresses raise a runtime error rather than falling back to an old detection address. This avoids intentionally using an obsolete observed address to block a different subscriber after dynamic reassignment.

block_subscriber first attempts legacy IP-owned rule cleanup, then resolves the current PPP address. It inspects subscriber-owned rules, removes those targeting obsolete source addresses, and distinguishes enabled and disabled rules for the current address. An enabled current rule is retained. Disabled current rules are re-enabled. If none exists, it creates a forward-chain source-address drop rule. It then clears matching tracked connections. The implementation does not move a newly added rule to a guaranteed position ahead of preexisting accepts or FastTrack rules.

_known_ips_for_pppoe reads distinct historical src_ip values for the name from intrusion_logs. _remove_legacy_subscriber_rules uses these addresses to identify older comments containing the prefix and address. This migrates from IP ownership to subscriber ownership. Historical address reuse means an address-only legacy rule may be ambiguous across subscribers. Cleanup therefore has historical assumptions that should be understood before treating ownership as perfect.

unblock_subscriber removes legacy rules, retrieves all matching subscriber-owned rules, collects their source addresses, and attempts deletion of each identifiable rule. It also resolves the subscriber's current active address when possible, but an offline subscriber does not prevent removal of already known owned rules. It clears tracked connections for collected addresses. The ordinary current restoration path does not require SSH, Paramiko, a PPP session disconnect, or a separate restorer process.

mark_restore_synced writes the current database time into firewall_synced_at only for the selected incident while it remains resolved. However, the current unblock_subscriber path does not assert that all deletion helpers succeeded or re-query owned rules before this marker is written. _remove_firewall_rule returns false after exhausting its fallbacks, and callers do not consistently propagate that result. Accordingly, the synchronization timestamp records that the agent reached that step; it is not an independently verified guarantee that every firewall rule disappeared.

The service retains legacy block_ip and unblock_ip helpers plus direct RouterOS protocol support. RouterOsDirectClient implements variable-length word framing, sentences, socket reads, ordinary login, challenge-response fallback using MD5, and command dispatch with trap detection. _set_firewall_rule_enabled tries several identifier forms through direct API and library fallbacks. _remove_firewall_rule tries identifiers with and without the RouterOS asterisk, different parameter names, direct API removal, and a final command-style fallback. These compatibility methods support different RouterOS API representations.

clear_connections queries source and destination tracked connections and attempts to remove each unique identifier. clear_ppp_active can disconnect PPP sessions by address, but is not invoked by the current subscriber reconciliation flow. Old variable names such as dst_ip in legacy blocking methods refer to the address being blocked even though the created filter matches source address. Function names should not be used to infer destination-based blocking policy.

The browser may display blocked or resolved before the router reaches that state. Blocking is continuously reconciled, while restoration is treated as pending until the synchronization marker is set. If the agent is stopped, database statuses still change but router enforcement does not follow. If another incident for the same subscriber remains blocked, restoring one incident leaves that subscriber blocked by design. Operators should inspect both incident state and the actual owned router rules when verifying outcomes.
''')
section('8. Website pages, user interface, and browser behavior', '''
resources/views/welcome.blade.php is the public landing page and contains the login form. It includes project branding, descriptive website content, theme handling, smooth scrolling to login, and a no-account modal used by the /login fallback when a dedicated auth.login template is absent. The root route is named logout even though it only returns the welcome page. Actual logout is /logout with the name logout.action; maintainers must distinguish those names.

resources/views/admin/dashboard.blade.php renders summary counts, attack and benign series, attack-type breakdowns, destination rankings, and risk distribution from DashboardController. Time windows are all, 24h, 7d, and 30d. Attack rate is attack rows divided by all selected rows, multiplied by one hundred. The controller deliberately avoids loading a full raw log collection on the dashboard. Most metrics use intrusion_logs with a null legacy status.

DashboardController selects the earliest available timestamp for the all-time window, falling back to five years before the current time when empty. The 24-hour chart uses MySQL DATE_FORMAT to group hourly data; daily windows use DATE. The 24-hour display generates twenty-four buckets from the computed start. Seven-day and thirty-day charts generate fixed day counts. A final partial current bucket can be excluded from generated chart labels even when aggregate counters include it. All-time charts include only dates actually returned by the grouping query.

The dashboard refresh endpoint returns selectedWindowLabel, window, totalLogs, attackCount, benignCount, attackRate, chartData, attackByType, topAttackedIPs, and riskLevelDistribution. Browser code requests an initial refresh after DOM readiness and polls every fifteen seconds while visible. It uses a cache-busting query value and a no-store fetch policy and prevents overlapping refreshes. There is no WebSocket broadcast or server-sent event implementation in this path.

resources/views/admin/logs.blade.php lists raw detections with ten rows per page. LogController supports date, type, and free-text query parameters. It inspects existing columns to choose an attack-type field and search fields, preferring is_malicious, label, type, classification, verdict, then attack_type. Search candidates include PPPoE name, addresses, protocol, notes, payload, info, and ports when present. Sorting is newest created_at first. Optional notes and subscriber columns are rendered only when found in the schema.

The logs type filter is legacy-compatible rather than a strict filter on risk_level. With attack_type selected, benign rows commonly have a null family, so comparing attack_type to benign or normal does not necessarily select the same records as risk_level benign. The benign branch adds an ungrouped OR condition, which can weaken earlier predicates. These are documented implementation caveats rather than corrected in this documentation task.

resources/views/admin/manage-threats.blade.php displays incidents, not an arbitrary grouping of source-IP rows. The controller eager-loads latestLog and adapts it into a display object containing incident_id, PPPoE identity, attack family, event count, group status, and last-detected timestamp. The default status filter is unresolved, with blocked and resolved alternatives. Date filtering applies to incident last_detected_at. Each page contains ten incidents, newest activity first.

The management page shows source address, PPPoE user, destination, protocol, attack type, event count, risk label, status, and optional notes. Protocol numbers are mapped into human-readable names, including TCP, UDP, ICMP, IPv6-related values, GRE, ESP, AH, OSPF, and SCTP. Missing values use a No Value presentation. Block and Restore confirmation modals explain that enforcement is performed by the local agent. Missing subscriber identity prevents these normal actions and results in a PPPoE-unavailable or legacy-block message.

Administrators can seed two configured demonstration subscriber incidents by supplying valid destination IPs. The controller creates one unresolved incident and one attack log for each target with predefined traffic statistics and attack probability 0.99. This does not generate network packets or exercise model inference. The example data is inserted into the same tables and is not globally excluded from dashboard counts. A real subscriber with the same demo identity could later be affected by an administrator's block action.

resources/views/admin/threat-reports.blade.php presents time-window statistics and a paginated detection table. ThreatReportsController computes selected total rows, attack rows, high-risk attack rows with probability greater than 0.8, and latest row while requiring null legacy log status. The data endpoint returns up to two hundred matching rows. The CSV endpoint streams rows using a cursor and includes timestamp, addresses, risk level, attack probability, forward packet count, and byte rate.

The CSV export is not a PDF export, and it is not identical to the paginated report query: its default window is 24h, while the report index defaults to all; export does not apply the null-status filter used by index and data. The JSON search includes notes without checking whether that optional column exists. Since the inspected migrations do not create notes, a nonempty report search can fail on a schema matching only these migrations. Formula-like values are not specifically escaped for spreadsheet interpretation in the CSV writer.

resources/views/admin/system-status.blade.php renders counts, processing estimates, health indicators, and a twenty-four-hour activity timeline. It refreshes the status alert fragment every fifteen seconds using /system-status/alert. The partial maps Healthy to success styling, Caution to warning, and Warning or Critical to danger. The page does not open a live RouterOS health connection, inspect the Python processes, or issue a detector probe; its indicators are derived from stored logs.

SystemStatusController counts current day, week, and month observations; computes last-hour and last-day rates; averages forward packet count and flow_bytes_s; and calculates a probability-based health label. Critical means more than one hundred probability-at-least-0.8 rows. Otherwise Warning means more than fifty such rows or an attack rate above thirty percent. Otherwise Caution means more than twenty such rows or attack rate above fifteen percent. Otherwise Healthy is used. These indicators are mostly all-time null-status metrics, even when a selected window constrains other page counts.

The status-page uptime value is based on timestamps and is not operating-system uptime or uninterrupted service availability. In selected-window mode, the current code can treat a Carbon start-window object like a log object when accessing created_at, leaving a null timestamp and producing a failure. This path has not been validated against the live deployment. Labels describing average bytes per flow actually use the stored byte-rate field. Probability categories are dashboard conventions separate from the model's calibrated classification threshold.

resources/views/admin/live-monitor.blade.php is a demonstration page. LiveMonitorController supplies placeholder status and a current timestamp. Browser code generates random initial values and random subsequent chart changes, updating flows, packet-rate-like values, alert-like values, and timestamps. These figures are not evidence of current network activity. They should be labeled or replaced before being used in an operational report as live measurements.

Shared visual structure is repeated across the administration templates rather than centralized in a single layout file. The templates use the static Sneat administration assets, Bootstrap styling, Boxicons, ApexCharts, and shared navigation links. The public landing folder also contains a separate HTML template with Bootstrap, animations, sliders, lightboxes, counters, and PHPMailer source. Their presence does not prove that all template demonstration pages or email handlers are active Laravel features.
''')
section('9. Authentication, authorization, and account lifecycle', '''
The actual website login uses App/Models/UsersAccount mapped explicitly to users_account. LoginController validates a syntactically valid email and a required string password, finds the first matching account, and verifies the stored password hash. A successful login regenerates the session identifier and stores user, user_id, and login_time, then redirects to the dashboard. Failure redirects to the login route with an invalid-credentials message. There is no registration route or account-management interface in the inspected route file.

UsersAccount permits fullname, age, sex, email, password, and role for mass assignment. Password is hidden from model serialization. Age is cast to an integer. The name accessor maps the expected display name to fullname, allowing templates to display it. canManageThreats returns true only when role is exactly admin. The role migration adds a string column with admin as its default so existing administrator accounts retain access.

CheckSession first requires a stored user object. For every protected request it reloads the corresponding UsersAccount row, invalidates the session when the account no longer exists, and replaces the stored object with the fresh row. Role changes therefore take effect without waiting for a new login. For unsafe HTTP methods it requires canManageThreats; otherwise it aborts with 403 and a view-only message. The route group currently includes all four threat mutation endpoints.

The management Blade template separately checks canManageThreats before rendering the Actions column, per-incident action cells and modals, the demo-seed button, and the demo-seed modal. This removes the controls from viewer HTML, while middleware rejects direct crafted POST requests. Hiding a button alone would not establish authorization; both checks are used. Viewer GET requests remain permitted, including viewing blocked and resolved records, filtering, JSON refresh, and report export.

ViewOnlyUserSeeder uses updateOrCreate for user@idms.com, assigns the viewer role, sets the display name to View Only User, and hashes the requested initial password. The password itself is intentionally absent from this documentation. Re-running the seeder resets that account's password and role. DatabaseSeeder calls it and also creates the framework Test User in the separate users table. Repeated generic DatabaseSeeder runs can conflict with the unique test email; use the dedicated viewer seeder for this task.

App/Models/User is the standard Laravel Authenticatable model using the users table, with notifications, factory support, datetime verification casting, and hashed password casting. config/auth.php still points its standard provider at this model. Custom UsersAccount session login does not automatically authenticate the standard guard or establish a Sanctum token. The protected /api/user route therefore has different authentication semantics from the website pages.

Logout flushes and invalidates the session and regenerates the CSRF token. The source exposes it as GET /logout. Protected forms include CSRF tokens through the web middleware stack. The source does not configure login throttling, multifactor login, email verification enforcement, a user-account password reset flow, or a record of which administrator requested each incident transition. The account password column is not automatically hashed by UsersAccount itself; callers other than the seeder must explicitly hash new passwords.

Session configuration defaults to database storage and a ten-minute idle lifetime unless environment settings override it. Automatic browser polling may refresh request activity, depending on session handling and whether the page remains visible. The public API ingestion endpoint is outside CheckSession and outside the viewer restriction. A viewer role should therefore not be represented as a complete machine-API access-control policy.
''')
section('10. Endpoint register and request responsibilities', '''
Website GET / returns the welcome page and is named logout. GET /login returns auth.login when that view exists, otherwise the welcome page with the missing-account modal flag. POST /login invokes LoginController login. GET /logout invokes LoginController logout and is named logout.action. These routes are outside the custom protected group.

Protected GET /dashboard invokes DashboardController gotoDashboard and accepts window. GET /dashboard/refresh invokes refreshData and returns JSON aggregates for window. GET /logs invokes LogController gotoLogs and accepts date, type, q, and pagination. GET /system-status invokes SystemStatusController gotoStatus and accepts window. GET /system-status/alert returns an HTML fragment from alertFragment. GET /live renders the demonstration monitor.

Protected GET /threat-reports invokes index and accepts window and pagination. GET /threat-reports/export invokes export and returns a streamed CSV. GET /threat-reports/data invokes data, accepts window and q, and returns an object with rows. These read endpoints are available to viewers as currently implemented.

Protected GET /manage-threats invokes ManageThreatController index and accepts status, date, and pagination. POST /threats/demo invokes demo and requires valid dst_ip_jimson and dst_ip_jimson2 values. POST /threats/{id}/block invokes block for an existing incident. POST /threats/{id}/ignore invokes ignore for restoration. POST /threats/bulk-action invokes bulkAction with action and ids. All four POST operations require the admin role through CheckSession.

The block and restore methods return 404 when the incident does not exist and 422 when its PPPoE identity is blank. They update incident state and redirect back with feedback. They do not call RouterOS directly. bulkAction treats the literal block action as blocked and any other nonempty action as resolved; it does not validate an explicit action enumeration. It filters out null subscriber names but not every possible blank string, and its success count is the number of submitted IDs rather than the actual updated-row count.

Laravel API POST /api/intrusion-log invokes IntrusionLogController store. It passes the entire request array into IntrusionLog create, with Eloquent fillable fields determining persisted attributes, and responds with status saved. There is no explicit request validation, API-key verification, authentication middleware, incident assignment, or custom error response in this controller. Extra fields are not automatically preserved if absent from fillable. Records inserted through this endpoint can lack an incident link, so they may appear in log statistics without appearing in incident management.

Laravel API GET /api/user returns the standard authenticated request user and uses auth:sanctum. The API route prefix is registered through bootstrap/app.php. GET /up is the framework health route registered in that bootstrap file; no custom detector, router, or MySQL probe is defined there by the project. Console routes contain only the scaffold inspire command, rather than an IDS scheduler.

Current Flask GET / returns a running message. GET /api/features returns count and ordered features. POST /api/analyze accepts the numerical feature payload and context, returns the classification fields described earlier, and attempts direct MySQL persistence. It returns 422 for missing expected features. No batch endpoint is defined. The Flask route root and Laravel root are different services even when both use an /api path segment.

The ZIP archive setup/ids_ensemble-api-main.zip contains a historical Flask app.py with the same three route shapes but different persistence and feature behavior. It maps compact keys to older CIC-style names with spaces and capitalization, fills missing model inputs with zero rather than rejecting the request, logs the prediction, and forwards it to LARAVEL_LOG_URL using requests. Optional LARAVEL_API_KEY is added as an X-API-KEY header, with LARAVEL_TIMEOUT defaulting to eight seconds. The current Laravel ingestion route does not verify that header. The archive's behavior should not be assumed to provide incident grouping or the fourteen-feature deployed contract.

The archived model file is only 134 bytes, consistent with a Git LFS pointer placeholder rather than a complete trained estimator. A deployment using that archive requires checking its actual model artifact. The archive requirements include gunicorn and an older scikit-learn pin, but the archive does not supply proof of a current hosting configuration or production process command. Its existence explains the Render-oriented comments and default collector URL that coexist with the newer local direct-database service.
''')
section('11. Database architecture, ownership, and relationships', '''
The operational database is shared between Laravel and Python. Laravel migrations define schema and Eloquent models handle website queries and incident actions. The detector uses PyMySQL transactions to insert observations and assign incidents. The blocker uses PyMySQL to read desired incident status and mark restoration synchronization. Website, detector, and blocker must point to the same intended database and compatible schema. Their environment files are separate; matching names in source do not enforce matching live settings.

intrusion_logs stores per-analysis observations. Its id is an auto-incrementing primary key. incident_id is a nullable indexed unsigned big integer added by the incident migration. It is an application-level relationship; the migration does not define a foreign-key constraint. pppoe_name is a nullable indexed string of length one hundred. src_ip and dst_ip are nullable strings of length forty-five, allowing textual IPv6 storage even though the current collector parses IPv4 only. Ports and protocol are nullable integers.

Flow measurements comprise nullable doubles for flow_duration, flow_pkts_s, flow_bytes_s, fwd_pkt_len_mean, bwd_pkt_len_mean, fwd_iat_mean, and bwd_iat_mean; nullable integers for tot_fwd_pkts, tot_bwd_pkts, tot_fwd_bytes, and tot_bwd_bytes. risk_level is a required string of length twenty. prob_attack is a required double. attack_type is a nullable string of length one hundred. created_at and updated_at are framework timestamp columns. A nullable legacy status enum is added and modified by later migrations.

The current final legacy status enum accepts blocked and resolved, with null representing unmarked observations. An intermediate migration also allowed pending, followed by another migration reclassifying pending rows as blocked and narrowing the enum. Current incident actions update threat_incidents instead of intrusion_logs.status. Therefore a blocked incident can have observation rows whose legacy status remains null, and those observations continue contributing to null-status dashboard queries. Distinguish legacy row visibility from current incident enforcement state.

threat_incidents stores the management unit. id is its auto-incrementing primary key. pppoe_name is a nullable indexed string of length one hundred. attack_type is a nullable indexed string. status is an indexed string of length twenty, default unresolved. event_count is an unsigned integer default one. first_detected_at is the first associated observation time; last_detected_at is the most recent and is indexed. blocked_at records a block request; restored_at records a restore request; firewall_synced_at records restoration synchronization processing. All three are nullable timestamps, as are the first and last detection times. created_at and updated_at track record lifecycle.

ThreatIncident logs is a has-many relationship using incident_id. latestLog is a has-one latestOfMany relationship selecting the latest related row by the model's default aggregation identity, not explicitly by a custom created_at ordering. IntrusionLog incident is a belongs-to relationship. No PPP subscriber master table, router-rule table, explicit user-to-incident action foreign key, or normalized attack-family reference table is defined. Subscriber identity is duplicated as text in incident and detection rows.

users_account is the actual custom website account table. It contains id, fullname, nullable unsigned tiny-integer age, nullable sex enum male/female/other, unique email, hashed password text storage, timestamps, and role added later. The role column is a string default admin, not a database enum restricted to only admin and viewer. Application permission checks accept only admin for mutation; any other role value is effectively read-only in this check.

The standard users table is separate and contains id, name, unique email, nullable email_verified_at, password, remember_token, and timestamps. password_reset_tokens stores a primary email, token, and nullable creation timestamp. sessions stores a primary string id, nullable indexed user_id, nullable client IP and user agent, a long-text session payload, and indexed last_activity. The user_id column is declared as a foreignId type but no constrained relationship is added in that scaffold migration.

personal_access_tokens is a Sanctum-compatible polymorphic token table. Its fields include id, tokenable_type, tokenable_id, name, unique sixty-four-character token, nullable abilities, nullable last_used_at, nullable indexed expires_at, and timestamps. The table's presence does not mean the custom website login issues tokens. No project route for creating these tokens is present in the inspected API file.

cache contains a primary key string, medium-text value, and integer expiration. cache_locks contains a primary key string, owner, and integer expiration. jobs stores id, indexed queue, long-text payload, attempts, nullable reserved_at, available_at, and created_at integers. job_batches stores id, name, total/pending/failed counts, failed job identifiers, nullable options, nullable cancelled_at, created_at, and nullable finished_at. failed_jobs stores id, unique UUID, connection, queue, payload, exception, and failed_at defaulting to the current time. These support framework infrastructure, not the local polling loop.

Laravel's migrations bookkeeping table is managed by the framework, rather than explicitly created in these application migrations. It records which migrations ran and their batch numbers. The file inventory also includes an optional local SQLite file if present; its existence is not proof that production uses SQLite. Several migration and query statements are MySQL-specific, so switching the full application to SQLite requires compatibility work even though isolated tests use SQLite.
''')
section('12. Migration history and data transformation rules', '''
The scaffold migrations create users, password reset tokens, sessions, cache tables, and queue tables. December 2025 migrations create intrusion_logs, personal access tokens, users_account, and a nullable intrusion-log status enum. June 2026 migrations broaden the status enum temporarily, backfill timestamps, reclassify pending statuses, and change one demo destination address. August 2026 migrations introduce PPPoE identity and incident tracking. October 2026 adds account roles.

The timestamp backfill replaces missing created_at or updated_at using the other timestamp and then the current database time. It has no meaningful reverse operation. The pending-status reclassification changes existing data, so a down migration that reopens the enum cannot reconstruct which blocked rows were formerly pending. The demo destination migration changes every matching destination address, not only rows explicitly marked as demonstrations, and its reverse similarly changes every matching replacement address.

The PPPoE identity migration adds the nullable indexed name column only if absent, then backfills two known demonstration source addresses to their configured subscriber names. It does not discover names for all historical observations. The incident migration creates indexed incident columns and adds nullable indexed incident_id to intrusion_logs without a database foreign key.

For historical attack logs, the incident migration reads all matching rows into memory, orders by ID, and groups by lowercased trimmed subscriber name, attack family, and legacy status. Each status group becomes a separate incident. Its first and last timestamps are taken from the first and last ID-ordered logs, rather than computing strict minimum and maximum timestamps. event_count equals group size. Legacy blocked and resolved values are retained; other values become unresolved. Each original row is linked to its incident. This historical grouping does not use the detector's fifteen-minute live grouping window.

Rolling back the incident migration drops the links and incident table, losing the new incident-management state. Applying migrations on a live database should be preceded by an appropriate backup and review of pending migrations. The targeted role migration used during the hosting walkthrough changes only users_account role, preserving existing account behavior with its admin default. The viewer account is created by a separate seeder; adding the role column alone does not create it.
''')
section('13. Machine-learning training, artifacts, and recorded results', '''
The maintained deployment training path is setup/Training/train_netflow_model.py with the contract in training_schema.py. The training README explicitly says it replaces the defective notebook path for deployment. Historical notebooks and manuscripts describe stacking ensembles and different thresholds; the current maintained script uses HistGradientBoostingClassifier for Stage A and RandomForestClassifier for Stage B. The project can still be described as a two-stage model, but the exact active estimators should be named accurately.

The dataset loader scans CSV files in the selected dataset directory, normalizes column names with Unicode normalization and non-alphanumeric replacement, locates a label-like column, and requires every feature in NETFLOW_FEATURES. It retains the fourteen numerical features, a normalized attack-family label, a capture filename, and a binary label determined by whether the normalized family is BENIGN. Numeric failures become missing values and infinities are replaced with missing values.

Optional max-rows-per-file sampling processes CSV chunks of fifty thousand rows, assigns reproducible random keys, and retains the smallest keys as a uniform sample across the file rather than simply taking the first records. The random seed begins at forty-two and varies by file index. The full dataset path loads all records. The source requires both binary classes and at least two attack families before training.

The split is seventy percent training, fifteen percent validation, and fifteen percent test, formed by a thirty-percent holdout followed by a half split. Stratification is on the binary label. It is not a capture-disjoint or time-disjoint split, even though capture filenames are retained. Related traffic from the same capture can therefore appear in more than one partition. This limits claims about generalization to unseen networks or future capture periods.

Each stage has its own cloned preprocessing pipeline: median imputation, StandardScaler, and a ColumnTransformer selecting the explicit feature list and dropping other columns. Independent cloning prevents Stage B from overwriting Stage A preprocessing state. Stage A uses two hundred maximum histogram-gradient-boosting iterations, learning rate 0.08, maximum thirty-one leaf nodes, balanced class weights, and random state forty-two. Stage B uses two hundred random-forest trees, maximum depth twenty-four, balanced-subsample weighting, all available jobs, and the same random state.

Stage A trains on all training rows and binary labels. Stage B trains only on training rows labeled attack and their family labels. select_threshold computes the precision-recall curve on validation probabilities, finds thresholds meeting target recall, and chooses the highest precision among those candidates. The default target recall is 0.95; invalid target values are rejected by argument parsing. The threshold is not chosen on the final test partition in the maintained script.

The bundle stores bin, mul, features, best_threshold, schema_version, dataset_profile, trained_at, and config. Configuration includes run type, sampling limit, random state, target recall, row count, and training elapsed time. dataset_profile is cicids-bootstrap. Full training defaults to setup/data/model/netflow_ids.joblib; a sampled run defaults to netflow_ids.smoke.joblib so it does not overwrite the full artifact. An explicit output argument can change either destination.

The trainer saves and reloads the artifact before evaluation and verifies that Stage A and Stage B do not share the same fitted preprocessor object. It computes full-test predictions using the stored threshold and writes a neighboring metrics JSON. A reproducible balanced check selects up to one hundred benign and one hundred attack rows from the untouched test partition. Small smoke datasets are marked incomplete when they cannot supply two hundred rows. The animated activity bar reports activity and elapsed time, not a trustworthy completion percentage.

The checked-in netflow_ids.metrics.json records a full run with schema mikrotik-netflow-v1 and threshold approximately 0.9865455954. Its full binary test matrix records 340751 true benign, 214 false-positive benign, 3537 missed attacks, and 80110 detected attacks across 424612 rows. Accuracy is approximately 99.1166 percent. Attack precision is approximately 99.7336 percent, attack recall 95.7715 percent, and attack F1 97.7124 percent. These are stored historical evaluation results, not measurements rerun for this documentation.

The balanced check records one hundred benign and one hundred attack rows, with one hundred benign correctly classified and ninety-seven attacks detected, giving 98.5 percent binary accuracy. Its attack-family report lists DDOS, DOS GOLDENEYE, DOS HULK, PORTSCAN, and SSH-PATATOR in that sample, with perfect family predictions on the selected true-attack rows. Family evaluation on true attacks is separate from end-to-end gated performance; missed attacks do not become operational family alerts even when Stage B would label them correctly in isolation.

The metrics file explicitly warns that CICIDS bootstrap results are not production validation and calls for labeled captures from the actual MikroTik environment. The old manuscripts' claims such as 99.92 percent binary accuracy or 99.83 percent multiclass accuracy must be attributed to those manuscripts and not substituted for the maintained artifact's metrics. Neither a metrics filename nor a repository copy proves that the live detector loaded that exact artifact or threshold.

The Dataset folder contains eight capture CSV files covering Monday working hours, Tuesday, Wednesday, Thursday morning web attacks, Thursday afternoon infiltration, Friday morning, Friday afternoon port scans, and a DDoS dataset. They are input material, not application runtime tables. Their individual names and sizes are included in the inventory. The full CSV payloads and serialized estimator internals are not reproduced in a prose reference.
''')
section('14. Configuration, dependencies, and file responsibilities', '''
Laravel requires PHP version 8.2 or newer under the composer constraint and Laravel framework major version twelve. Sanctum supports token authentication infrastructure; Tinker provides interactive application inspection. Development dependencies include Faker, Pail, Pint, Sail, Mockery, Collision, Pest, and the Pest Laravel plugin. composer.lock records resolved package versions and should be distinguished from composer.json's allowed ranges. Installed vendor source is third-party framework code, not custom detection logic.

The frontend manifest includes Vite major seven, Laravel's Vite plugin, Tailwind major four and its Vite integration, Axios, and Concurrently. npm run dev starts Vite; npm run build produces deployment assets. vite.config.js defines resources/css/app.css and resources/js/app.js as entrypoints and ignores compiled Blade cache paths in its watcher. resources/js/bootstrap.js initializes Axios and its XMLHttpRequest header. The administration pages often use public static assets directly, so not every visible element depends on rebuilding Vite.

The Python requirements include Flask, Gunicorn, joblib, NumPy, pandas, scikit-learn pinned to 1.6.1, XGBoost, Requests, python-dotenv, PyMySQL, and routeros-api. Some retained libraries serve historical models or experiments rather than the maintained trainer's active estimator selection. Artifact compatibility depends on the library versions used for serialization and loading. The Windows launcher assumes python resolves to an environment containing the required packages and does not activate a named virtual environment.

bootstrap/app.php registers web, API, console routes, /up, and the check.session middleware alias. bootstrap/providers.php registers AppServiceProvider. The provider currently contains no active custom application boot behavior; a force-HTTPS example is commented out. public/index.php is the web front controller loading Composer and the application bootstrap. artisan is the command-line entrypoint used for migrations, seeders, caches, tests, and framework commands.

config/app.php defines application identity, environment, debug flag, URL, locale, encryption configuration, and a hardcoded UTC timezone in this snapshot. config/database.php defines database connections; config/session.php defines storage and cookies; config/cache.php defines caching; config/queue.php defines asynchronous job infrastructure; config/filesystems.php defines storage disks; config/logging.php defines application log channels; config/mail.php and config/services.php define external service configuration; config/auth.php defines standard authentication; config/sanctum.php defines Sanctum behavior. Environment variables can override the supported settings, but secret values are not part of this reference.

Python settings shared in setup/.env include DB_HOST, DB_PORT, DB_DATABASE, DB_USERNAME, DB_PASSWORD, DB_TABLE, DB_INCIDENT_TABLE, MIKROTIK_HOST, MIKROTIK_PORT, MIKROTIK_USER, MIKROTIK_PASSWORD, API_URL, IDS_MODEL_FILENAME, IDS_THRESHOLD, INCIDENT_WINDOW_MINUTES, BLOCK_SCAN_INTERVAL, BLOCK_RULE_PREFIX, and LOG_DIR where referenced. The environment variable register later lists source references and safe descriptions. NETFLOW_PORT, ANALYSIS_INTERVAL, and MIN_FLOWS_FOR_ANALYSIS are current script constants rather than environment-controlled settings.

Default API_URL in the collector points to a Render-hosted inference endpoint, while setup README describes a local detector. The actual selected service depends on setup/.env. No environment values are published here. This mismatch must be resolved operationally by checking the configured inference URL, feature endpoint, model schema, and persistence path. A running local detector alone does not prove the collector is sending to it.

The repository includes .gitignore entries for environment files, node_modules, vendor, generated frontend artifacts, selected model artifacts, the primary IDS virtual environment, and training Dataset. Some alternate model files and a second virtual environment may still exist locally. Git ignore rules control new tracking, not already tracked content or web-server exposure. A file under public can become downloadable after deployment; project documentation therefore omits private configuration and runtime personal records.

The two ngrok executables are optional tunneling tools by location; no main runtime script launches them. Public ZIP template bundles and framework demonstration HTML pages are distribution or design assets. A PHPMailer library in the landing template is not evidence of an implemented contact-form mail service. Compiled Python cache files, notebook checkpoints, Laravel caches, logs, sessions, and virtual environments are generated development or runtime artifacts. Their purpose is recorded without treating them as independent system features.
''')
section('15. Installation, startup, hosting, and change deployment', '''
Prepare a monitoring host with a compatible Python interpreter and install setup/data/requirements.txt into an appropriate environment. Place a trusted compatible model bundle in setup/data/model and configure setup/.env for router access, database connectivity, inference URL, and optional thresholds and polling intervals. Configure the router's Traffic Flow export to the monitoring host and verify that port 2055 traffic and RouterOS API access are reachable. The repository's current collector requires NetFlow version five.

Prepare Laravel using Composer dependencies, the application environment, an existing valid application key, and the intended database. Ensure storage and bootstrap/cache are writable by the web process. Point the site's document root to Laravel public where hosting permits. Build frontend source assets if the deployment uses their Vite outputs. Run reviewed migrations and only the intended seeders. Do not regenerate an established production application key as a routine deployment step, because it changes encryption behavior.

Start the detector before the collector when API_URL targets it, then start the collector and the blocker. setup/Start-IDS.bat performs this order using separate Windows command windows but does not wait for API readiness, validate connectivity, restart failed services, or shut them down as a supervised group. The manual run instructions name Detection Process.py, Mikrotek PPPoe Sniffer.py, and Pending Blocker.py. Keep the services running for continuous collection and enforcement.

Verify each boundary independently: the detector feature endpoint reports the expected contract; the router emits accepted NetFlow; the collector recognizes current PPPoE names; inference accepts the required features; detection rows appear in the intended database; attack rows link to incidents; Laravel displays those rows; an authorized block changes the incident; the agent resolves the current subscriber address; the correct owned rule appears; restoration removes the appropriate rules subject to other blocked incidents. Tests should use controlled authorized subscribers and distinguish demo insertions from real model outputs.

For the hosted viewer-role change, deploy LoginController, CheckSession, UsersAccount, the manage-threats Blade view, the role migration, and the dedicated viewer seeder, plus DatabaseSeeder if preserving the full repository's seed configuration. Apply the role migration, then execute ViewOnlyUserSeeder. The hosting walkthrough established a working account-relative project path of domains/dissertationsystem.online/public_html/intrusion-detection-monitoring-system/artisan in the PHP cron form. The form automatically prepends the hosting account's home path and PHP executable; entering the full home prefix again duplicates it.

The successful targeted command runs artisan migrate with the path database/migrations/2026_10_29_000001_add_role_to_users_account_table.php and the production force option. The next operation runs artisan db:seed with class ViewOnlyUserSeeder and the force option. Temporary cron scheduling used an asterisk in every field; a day-of-month value zero is invalid. Remove each temporary cron job after successful execution. Leaving the seeder scheduled repeatedly resets the viewer account credentials and role. A seeding output containing only an informational seeding message can be normal; verify successful login and role behavior separately.

The correct path was established through user-provided cron results, not by deriving a server filesystem path from the File Manager URL. A browser File Manager address is not an absolute shell path. Production credentials and host-specific PHP executables must be checked on the hosting account. This reference documents the demonstrated workflow without guaranteeing that later hosting plan changes or server moves preserve identical paths.

After Blade changes, clear stale compiled views through the normal Artisan deployment process if required. Cached routes or configuration must be refreshed appropriately when those files change. Do not use migrate:fresh or other destructive rebuild commands against the live operational database. A role migration rollback removes the role column while the new model and middleware expect it, so application and schema rollback must be coordinated. Restore access should be verified using the actual router, not solely website badges.
''')
section('16. Tests, verification boundaries, and acceptance scenarios', '''
tests/Feature/PppoeBlockingIdentityTest.php builds minimal in-memory tables and exercises ManageThreatController directly. Its first test ensures block and restore change only the selected incident, leave a newer same-subscriber incident unresolved, leave a different subscriber untouched, and do not overwrite individual log legacy statuses. Its second test verifies that a resolved historical incident and a new unresolved attack remain visible as separate management entries under their respective status filters.

tests/Feature/ViewOnlyAccessTest.php creates the custom account schema and role column, seeds the viewer, and verifies login with the requested initial credentials. It posts to demo, block, restore, and bulk-action routes and expects forbidden responses. It uses a temporary protected endpoint to verify administrators may submit and that demotion affects the next request. It verifies a viewer can read and that the management HTML omits the demo modal and Actions column, while the administrator sees the modal.

tests/Feature/ExampleTest.php checks the public page response and tests/Unit/ExampleTest.php provides a basic truth assertion. tests/TestCase.php supplies the framework base; tests/Pest.php binds feature tests to it. phpunit.xml sets a testing environment, in-memory SQLite, array cache and sessions, synchronous queues, and reduced bcrypt rounds. The feature tests intentionally construct their needed tables rather than running the entire migration chain, which contains MySQL-specific statements.

Earlier verification for the viewer change reported eight passing tests and twenty-eight assertions. Those tests demonstrate isolated login and authorization behavior, incident selection, and page fragments. They do not exercise real RouterOS APIs, NetFlow parsing, machine-learning inference, MySQL schema upgrades, hosted cron execution, external HTTP connectivity, or all dashboard/system-status controller branches. The documented test count must not be presented as end-to-end production certification.

An administrator acceptance test should establish a controlled subscriber identity and its current dynamic address, observe a real attack incident or clearly identified demonstration incident, submit a block, and verify the expected owned source-address rule. It should reconnect the subscriber with a changed address and confirm that reconciliation removes the stale owned rule and protects the intended identity. A restore test should verify deletion on the router, not only the resolved badge or synchronization marker.

A multiple-incident scenario should block two incidents for one name, restore only one, and verify that the remaining blocked incident prevents internet restoration. Restoring the final blocked incident should allow cleanup. A history scenario should resolve an incident, generate a later attack for the same name and family, and verify a new unresolved incident is created. A viewer scenario should verify read access to all intended pages and direct POST rejection for each mutation endpoint.

A persistence-failure scenario should simulate an unavailable database in a safe test environment and confirm that inference success can coexist with storage failure, as implemented. A collector-idle scenario should test the traffic-driven timer behavior. A schema-mismatch scenario should omit a required feature and expect a 422 from the maintained detector. A report-search scenario should exercise q on the production schema because the optional notes assumption is not covered by current tests.
''')
section('17. Operational limitations and accurate interpretation', '''
The system is flow-based and periodically aggregates per-subscriber records. It does not inspect full packet payloads, decrypt TLS, provide signature updates, reconstruct every session, or prove compromise from a probability alone. A subscriber aggregate can contain many unrelated conversations. Displayed destination metadata is context from one latest record rather than a complete explanation of the prediction. Exporter configuration, filtering, and feature approximations influence what is observed.

Network availability is a dependency in several directions: router to collector UDP, collector to inference HTTP, detector to MySQL, Laravel to MySQL, blocker to MySQL, and blocker to RouterOS. An online website does not imply those local links are healthy. Historical log statistics can remain visible when collection has stopped. The system-status and live-monitor pages do not provide reliable independent process-liveness checks in their current implementation.

The current Flask API lacks source authentication and runs its direct entrypoint with debug enabled on all interfaces. RouterOS pools use unencrypted API settings. The Laravel ingestion API accepts mass-assigned permitted fields without explicit validation or API authentication. These are concrete source observations. Deployment should establish appropriate access boundaries, transport protections, request validation, and production process configuration rather than assuming website roles cover machine interfaces.

Database failure is logged but not returned as an inference error. There is no durable store-and-forward buffer or end-to-end message identifier. Incident creation lacks a unique eligibility constraint. Timestamp sources differ: Python uses host-local datetime, Laravel config declares UTC, browser timestamps use browser locale, and database synchronization uses NOW. Without consistent timezone configuration and clock synchronization, time filters, incident windows, uptime estimates, and activity charts can be misleading.

The viewer role is enforced on unsafe requests inside the protected web group. It does not automatically cover future mutations added outside that group, future GET endpoints that change state, or the separate ingestion API. The account role defaults to admin for compatibility; newly created custom accounts should have an explicit intended role. The initial viewer password is a user-requested setup credential, not a recommended permanent access policy, and is omitted here.

Status has two meanings in the database: legacy observation status and current incident status. Dashboard null-status filters do not mean no currently blocked subscriber contributed those observations. Incident event_count summarizes linked detections and may diverge if logs are inserted through an alternate path or manipulated independently. No database foreign key prevents orphaned incident links. The optional notes field is referenced in some paths without a matching migration in this snapshot.

Restoration synchronization is not fully verified by the current subscriber cleanup path because deletion failures can be swallowed. Block requests are not marked with a corresponding success acknowledgement and continue reconciliation. A new drop rule may not be effective if earlier router rules accept or FastTrack the traffic. Removing an IDS-owned drop rule cannot override unrelated firewall rules, service outages, upstream failures, or PPP account restrictions; a resolved incident therefore does not guarantee working internet service.

Historical research describes broader claims and architectures, including packet-based capture, stacking ensembles, and high benchmark scores. Those materials preserve project rationale but do not override current code. Current maintained training is a NetFlow-compatible bootstrap pipeline with a documented distribution gap. Production validation should use labeled data from the actual environment and separately measure false positives, missed attacks, family confusion, decision latency, and operational impact.
''')
section('18. Maintenance, troubleshooting, and glossary', '''
When no detections appear, inspect the collector console and merged_ids.log for accepted NetFlow version-five records, verify the exporter sender matches MIKROTIK_HOST, check excluded address prefixes, verify PPP discovery, and confirm at least ten records exist for the selected client. Consider queue length and traffic-driven timers before assuming the classifier is failing. Then confirm the configured API_URL targets the intended current detector rather than an older hosted API.

When inference rejects a request, compare /api/features against the sender's normalized keys and the training schema. Missing-model-features errors indicate a contract mismatch; present but malformed values may instead be coerced to zero. When inference succeeds but the website is unchanged, check database-save errors and the intended database/table settings, since an HTTP success does not certify persistence. Compare observation timestamps and avoid using the synthetic live-monitor values as a diagnostic.

When a block badge appears but traffic continues, verify the local agent is running, inspect blocked_threat_sync.log, confirm the subscriber currently has an active PPP session, inspect owned firewall rules and their source address, review rule order and FastTrack, and inspect connectivity to both database and router. When restoration does not work, check for another blocked incident for the same name, failed removal fallbacks, unrelated router rules, and whether the synchronization marker was set despite residual rules.

Keep the classifier bundle, ordered features, schema version, training script, metrics file, and runtime dependency versions together as a release unit. Keep website code, Python writers, and database schema compatible during upgrades. Maintain backups for operational rows, account data, incident state, and trusted model artifacts; runtime prediction and enforcement logs support diagnosis but do not replace database backups. Log rotation and retention are not implemented in the Python file handlers.

PPPoE means Point-to-Point Protocol over Ethernet; its username is treated as the stable subscriber identity. RouterOS is the MikroTik router operating system and API target. NetFlow is exported traffic-summary telemetry. A flow record summarizes traffic counters and endpoints rather than holding packet content. Forward traffic is traffic where the selected subscriber is source; backward traffic is traffic where it is destination. IAT means inter-arrival time, although the collector's stored IAT fields are aggregate approximations.

Stage A is binary benign-versus-attack inference. Stage B is attack-family inference gated by Stage A. A threshold converts attack probability to a decision. Precision describes the proportion of predicted attacks that are truly attacks in an evaluated dataset; recall describes the proportion of actual attacks detected; F1 combines precision and recall. Bootstrap metrics are development evidence from a surrogate dataset, not deployed accuracy. An incident is a management record linking observations for a subscriber and family according to grouping rules.

Unresolved means awaiting an administrator decision. Blocked means the agent should maintain a subscriber-owned drop rule. Resolved means restoration has been requested and the incident remains in history. firewall_synced_at is the agent's restoration-processing marker with the verification limitations described earlier. CSRF protects browser state-changing requests against cross-site forgery within the web stack. Sanctum is the separate Laravel token/session authentication infrastructure, not the custom UsersAccount login itself.

The following automatically derived registers anchor the narrative to specific local source files and line locations. They are a reference map, not executable instructions or an attempt to repeat every dependency's source. Research extracts are explicitly historical and preserve their own claims for comparison. The final inventory lists project files while grouping installed dependencies, private runtime data, and generated caches by category.
''')


def read(path):
    return path.read_text(encoding='utf-8-sig', errors='replace')


def scrub(text):
    # Historical documents may contain addresses and access settings. Never publish private assignments.
    text = re.sub(r'(?im)^.*(?:password|secret|api[_ -]?key|auth[_ -]?token)\s*[:=].*$', '[Private credential assignment omitted.]', text)
    text = re.sub(r'(?i)\b(?:mysql|postgres|https?)://[^\s/:]+:[^\s@]+@[^\s]+', '[Credential-bearing connection URL omitted]', text)
    text = re.sub(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b', '[email address omitted]', text)
    return text


section('19. Source member register: Python services and training', '''
Each entry below names a source-defined class or function and its line location. A class entry groups related responsibilities; a function entry records its local documentation or main operation families. Call names are a source navigation aid, not a claim that every call executes on every request. The detailed behavioral explanations in the earlier chapters take precedence over short historical docstrings.
''')
py_paths = sorted((ROOT / 'setup').glob('*.py')) + sorted((ROOT / 'setup' / 'Training').glob('*.py'))
for path in py_paths:
    source = read(path)
    tree = ast.parse(source)
    paras = []
    def members(nodes, prefix=''):
        for node in nodes:
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                name = prefix + node.name
                doc = ast.get_docstring(node)
                desc = scrub(re.sub(r'\s+', ' ', doc or ''))
                if isinstance(node, ast.ClassDef):
                    paras.append(f'{name}, line {node.lineno}: class grouping {sum(isinstance(x, (ast.FunctionDef, ast.AsyncFunctionDef)) for x in node.body)} directly declared methods. ' + desc)
                    members(node.body, name + '.')
                else:
                    calls=[]
                    for call in ast.walk(node):
                        if isinstance(call, ast.Call):
                            fun=call.func
                            n=fun.id if isinstance(fun,ast.Name) else fun.attr if isinstance(fun,ast.Attribute) else ''
                            if n and n not in calls and n not in ['print','str','int','float','len','get','append','info','debug','warning','error','sum','max','min','isinstance','range','strip','lower','join','items','values','keys']:
                                calls.append(n)
                    args=', '.join(a.arg for a in node.args.args if a.arg not in ['self','cls']) or 'no explicit positional parameters'
                    paras.append(f'{name}, line {node.lineno}: accepts {args}. ' + (desc+' ' if desc else '') + ('Referenced operations include '+', '.join(calls)+'.' if calls else 'A local helper providing the return or state update described by its name and enclosing component.'))
    members(tree.body)
    section(path.relative_to(ROOT).as_posix(), '\n\n'.join(paras) or 'Declares the ordered NetFlow feature contract and schema identifier used by the maintained training path.',2)

section('20. Source member register: PHP controllers, models, and tests', '''
This register lists every declared PHP method in the custom application and test files, together with its source line. Earlier chapters explain runtime behavior; these entries make source navigation and maintenance assignment precise. An endpoint is only active when its route registration and middleware permit it; a helper's existence does not establish a public endpoint.
''')
for folder in ['app', 'tests', 'database/seeders']:
    for path in sorted((ROOT/folder).rglob('*.php')):
        source=read(path)
        entries=[]
        for match in re.finditer(r'(?:public|protected|private)\s+(?:static\s+)?function\s+(\w+)\s*\(([^)]*)\)',source):
            line=source[:match.start()].count('\n')+1
            name=match.group(1)
            summaries={
                'login':'Validates email and password, verifies the custom account hash, and establishes a regenerated website session.',
                'logout':'Flushes and invalidates the session, regenerates the CSRF token, and redirects to login.',
                'handle':'Reloads the custom account and enforces session existence and admin-only unsafe requests.',
                'canManageThreats':'Returns whether the role is exactly admin.',
                'index':'Builds the page query, selected filters, statistics or incident display data, and returns the associated view.',
                'block':'Marks one identified subscriber incident blocked and resets restoration synchronization state.',
                'ignore':'Marks one identified subscriber incident resolved and pending router restoration processing.',
                'bulkAction':'Applies a submitted block or resolution action to selected incident IDs.',
                'demo':'Validates two destination addresses and creates demonstration incidents and observation rows.',
                'store':'Creates one mass-assignable intrusion log from request data and returns a saved status.',
                'export':'Streams a CSV using a database cursor and the selected time window.',
                'data':'Returns up to two hundred filtered report observations in a JSON rows object.',
                'run':'Executes the seeder account creation or update operations.',
                'logs':'Defines the incident-to-observations has-many relation.',
                'latestLog':'Defines a latest-of-many related observation for management display.',
                'incident':'Defines the observation-to-incident belongs-to relation.',
                'casts':'Defines attribute conversions used by Eloquent.',
                'getNameAttribute':'Maps display name to fullname for the custom account.',
                'refreshData':'Returns selected dashboard metrics for browser polling.',
                'gotoDashboard':'Builds initial dashboard metrics and time-window labels.',
                'gotoLogs':'Applies optional log filters, schema-compatible searches, and ten-row pagination.',
                'gotoLive':'Returns placeholder data for the simulated live-monitor page.',
                'gotoStatus':'Builds system statistics, processing metrics, health labels, and the activity timeline.',
                'alertFragment':'Renders only the database-derived health alert for AJAX polling.',
                'setUp':'Builds the isolated test preconditions and minimal database schema.',
                'register':'Framework service registration hook; no custom behavior in this provider.',
                'boot':'Framework provider startup hook; no active custom HTTPS forcing in this snapshot.',
            }
            desc=summaries.get(name)
            if not desc:
                readable=re.sub(r'(?<!^)(?=[A-Z])',' ',name).replace('_',' ').strip()
                desc=('Regression test for '+readable.removeprefix('test ') if name.startswith('test') else 'Internal component helper for '+readable)+'.'
            entries.append(f'{name}, line {line}. {desc}')
        if entries: section(path.relative_to(ROOT).as_posix(),'\n\n'.join(entries),2)

section('21. Exact schema declaration register', '''
The following declarations are transcribed automatically into prose from every application migration. They supplement the consolidated database dictionary with field sizes, explicit defaults, indexes, and rollback history. Entries from later migrations may alter earlier declarations; they must be read in filename order. The register describes declared schema, not a queried production schema. Row transformations are explained in the migration-history chapter.
''')
for path in sorted((ROOT/'database/migrations').glob('*.php')):
    source=read(path)
    paras=[]
    for match in re.finditer(r"Schema::(create|table|dropIfExists)\('([^']+)'",source):
        paras.append(f'Schema operation {match.group(1)} concerns table {match.group(2)} at line {source[:match.start()].count(chr(10))+1}.')
    for match in re.finditer(r'\$table->([A-Za-z]+)\(([^;\n]*)\);',source):
        typ,tail=match.groups()
        # This is explanatory text, not code rendering; preserve field names and type parameters.
        parts=tail.split(')->')
        params=parts[0].replace("'",'').replace('"','').replace(')','').strip()
        modifiers='; '.join(x.replace('(',' ').replace(')','').replace("'",'') for x in parts[1:])
        paras.append(f'Declaration {typ}: {params or "framework default columns"}.' + (f' Modifiers: {modifiers}.' if modifiers else ''))
    for match in re.finditer(r'(?:ALTER TABLE|UPDATE)\s+(\w+)[^";]*',source,re.I):
        t=re.sub(r'\s+',' ',match.group(0)).strip()
        if len(t)<420: paras.append('SQL migration operation: '+t+'.')
    section(path.relative_to(ROOT).as_posix(),'\n\n'.join(paras),2)

section('22. Environment variable register without private values', '''
This list is derived from getenv and env references in application configuration, bootstrap, and current setup scripts. Values in private environment files are not read into the document. Variables marked as access material must be configured privately. A variable referenced only in a historical archive belongs to that archive's deployment, not automatically the active detector.
''')
env_refs={}
for base,pattern in [('config','*.php'),('app','*.php'),('bootstrap','*.php'),('setup','*.py')]:
    paths=(ROOT/base).rglob(pattern) if base!='setup' else (ROOT/base).glob(pattern)
    for path in paths:
        for name in re.findall(r'(?:os\.getenv|env)\(\s*[\'"]([A-Z][A-Z0-9_]*)[\'"]',read(path)):
            env_refs.setdefault(name,set()).add(path.relative_to(ROOT).as_posix())
descs={
    'API_URL':'Target inference request URL; determines whether the collector uses local or hosted inference.',
    'IDS_MODEL_FILENAME':'Trusted model bundle filename under setup/data/model.',
    'IDS_THRESHOLD':'Optional override of the stored binary classification threshold.',
    'INCIDENT_WINDOW_MINUTES':'Eligibility window for reusing unresolved incidents; default fifteen minutes.',
    'BLOCK_SCAN_INTERVAL':'Seconds between enforcement scans; default twenty.',
    'BLOCK_RULE_PREFIX':'Ownership prefix used in managed router comments.',
    'DB_TABLE':'Observation table name for Python persistence and lookup.',
    'DB_INCIDENT_TABLE':'Incident-management table name for Python grouping and enforcement.',
    'LOG_DIR':'Prediction log output directory for the detector.',
    'MIKROTIK_HOST':'Router address and required NetFlow sender address.',
    'MIKROTIK_PORT':'RouterOS API port; default 8728 in the scripts.',
}
paras=[]
for name,refs in sorted(env_refs.items()):
    desc=descs.get(name)
    if not desc:
        if any(s in name for s in ['PASSWORD','SECRET','TOKEN','KEY']): desc='Private access or cryptographic configuration; do not publish its configured value.'
        elif name.startswith('DB_'): desc='Database connection, storage, or transport configuration.'
        elif name.startswith('SESSION_'): desc='Website session persistence, lifetime, or cookie policy configuration.'
        elif name.startswith(('MAIL_','AWS_','POSTMARK_','RESEND_','SLACK_')): desc='Optional framework mail, cloud, or external service configuration; not proof that the service is used by the IDS path.'
        elif name.startswith(('REDIS_','MEMCACHED_','CACHE_')): desc='Framework cache or backing-store configuration.'
        elif name.startswith(('QUEUE_','SQS_')): desc='Framework asynchronous queue configuration; separate from the Python polling service.'
        elif name.startswith('LOG_'): desc='Laravel logging channel, level, or transport configuration.'
        else: desc='Framework or component setting; its exact consumers are the referenced source files.'
    paras.append(name+'. '+desc+' Referenced by '+', '.join(sorted(refs))+'.')
section('Configuration names and consumers','\n\n'.join(paras),2)

section('23. Notebook and archived API evidence register', '''
Notebook contents are development evidence, not automatically executed runtime components. Code cells are summarized as prose; stored outputs are not used to certify live behavior. Empty notebooks and checkpoints are identified as such. The historical API ZIP is inspected as text without executing it or loading its binary artifact.
''')
for path in sorted((ROOT/'setup/Training').rglob('*.ipynb')):
    notebook=json.loads(read(path));cells=notebook.get('cells',[])
    code='\n'.join(''.join(c.get('source',[])) for c in cells if c.get('cell_type')=='code')
    names=re.findall(r'\b(?:class|def)\s+(\w+)',code)
    facts=[f'The notebook contains {len(cells)} cells and {sum(c.get("cell_type")=="code" for c in cells)} code cells.']
    if not code.strip(): facts.append('No substantive code source is stored in this notebook snapshot.')
    if names: facts.append('Defined helpers or classes: '+', '.join(names)+'.')
    if 'StackingClassifier' in code: facts.append('Historical training experiment using a stacking classifier with random forest, gradient boosting, optional XGBoost, and a logistic-regression meta-model. It is distinct from the maintained NetFlow trainer.')
    if 'PRED_THR' in code: facts.append('Contains a fixed prediction-threshold experiment at 0.80 and writes a predicted-attacks CSV; this is not the maintained artifact threshold.')
    if 'SUSP_THR' in code: facts.append('Contains historical watch and suspicious severity gates, which are not implemented as those categories in the current Flask detector.')
    if 'BOT' in code: facts.append('Contains an offline BOT-label prediction or probability inspection experiment, not an active network attack generator.')
    if 'predicted_attacks_only.csv' in code: facts.append('Runs inference on CSV rows, retains predicted attacks, adds predicted family and optional family confidence, and exports a results CSV. The filename Network Attack Simulation does not establish actual packet generation.')
    if 'pre' in code and 'StackingClassifier' in code: facts.append('The historical notebook reuses a preprocessor object across stage definitions; the maintained trainer deliberately clones preprocessing state and evaluates after artifact reload.')
    section(path.relative_to(ROOT).as_posix(),'\n\n'.join(facts),2)
with zipfile.ZipFile(ROOT/'setup/ids_ensemble-api-main.zip') as archive:
    facts=[]
    for info in archive.infolist():
        if not info.is_dir(): facts.append(f'{info.filename}: {info.file_size} bytes. '+ ('Model placeholder, not loaded for documentation.' if info.filename.endswith('.joblib') else 'Historical API source or dependency/distribution metadata.'))
    section('Archived API member inventory','\n\n'.join(facts),2)

section('24. Historical research manuscripts: text extracts', '''
The following extracts preserve text from the project-owned Word manuscripts found in setup/Training so the documentation includes the original rationale, literature discussion, methodology ideas, and evaluation claims as well as the current implementation. Their claims are historical manuscript claims and have not been independently verified. They can conflict with active source and checked-in metrics. The current-source chapters must be used to describe what this project revision actually implements. Private email addresses and credential assignments are omitted. Tables are flattened into sentences so this remains a text-only reference; embedded figures and images are not reproduced.
''')
for path in sorted((ROOT/'setup/Training').glob('*.docx')):
    doc=Document(path)
    paragraphs=[scrub(p.text.strip()) for p in doc.paragraphs if p.text.strip()]
    paragraphs.insert(0,'Historical source manuscript: '+path.relative_to(ROOT).as_posix()+'. This appendix preserves source text and does not endorse its numerical results or treat its architecture as the current deployment.')
    for index,table in enumerate(doc.tables,1):
        paragraphs.append(f'Manuscript table {index}, flattened text follows.')
        for row in table.rows:
            cells=list(dict.fromkeys(scrub(c.text.strip().replace('\n','; ')) for c in row.cells))
            paragraphs.append(' ; '.join(cells))
    section(path.name,'\n\n'.join(paragraphs),2)

section('25. Complete project-file inventory and uses', '''
This inventory covers first-party source, configuration, tests, static website assets, setup material, datasets, archives, and generated artifacts visible in the project. Each file is named with its purpose and byte size. Installed PHP dependencies, JavaScript dependencies, Git internals, virtual environments, private environment files, and private runtime session/log contents are grouped rather than expanded. These are reproducible dependency or sensitive runtime stores, not hundreds of additional custom IDS modules. Inventory sizes describe this local snapshot only.
''')
excluded_dirs={'.git','vendor','node_modules','.venv-ids','.venv-model-test','.idea','.vscode','.phpunit.cache'}
private_or_generated=[]
files=[]
def collect(folder):
    for path in sorted(folder.iterdir()):
        rel=path.relative_to(ROOT)
        if path.is_dir():
            if path.name in excluded_dirs:
                private_or_generated.append(rel.as_posix()+': installed dependencies, editor/test state, virtual environment, or version-control metadata; internal files not expanded.')
            elif rel.as_posix() in ['storage','bootstrap/cache']:
                private_or_generated.append(rel.as_posix()+': generated framework runtime storage, caches, sessions, views, and logs; contents not published.')
            elif rel.as_posix()=='public/about_system':
                continue
            else: collect(path)
        elif path.name.startswith('.env') or path.name in ['auth.json','commands.txt']:
            private_or_generated.append(rel.as_posix()+': private environment or local operational file; values and contents not published.')
        else: files.append(path)
collect(ROOT)
purpose_map={
    '.gitignore':'Version-control exclusion rules for private and generated files.',
    '.gitattributes':'Version-control line-ending or file handling attributes.',
    '.editorconfig':'Shared editor formatting settings.',
    'README.md':'Framework scaffold overview; operational IDS setup is documented separately under setup.',
    'artisan':'Laravel command-line bootstrap for migrations, seeders, caches, and other framework commands.',
    'composer.json':'PHP dependency ranges, autoload mappings, and development/setup commands.',
    'composer.lock':'Resolved PHP dependency versions for reproducible installs.',
    'package.json':'Frontend build tooling and scripts.',
    'package-lock.json':'Resolved npm dependencies and integrity metadata.',
    'phpunit.xml':'Isolated PHP test environment configuration.',
    'vite.config.js':'Frontend build entrypoints and development watcher configuration.',
    'public/index.php':'Laravel public HTTP front controller.',
    'public/robots.txt':'Crawler instructions; not an authentication control.',
    'setup/README.md':'Local services, startup sequence, shared database, and incident lifecycle guide.',
    'setup/Start-IDS.bat':'Windows launcher opening the three current Python services in separate consoles.',
    'setup/Detection Process.py':'Current Flask two-stage inference API and direct transactional MySQL observation/incident writer.',
    'setup/Mikrotek PPPoe Sniffer.py':'NetFlow v5 UDP collector, PPP identity discovery, numerical aggregation, and inference client.',
    'setup/Pending Blocker.py':'Database-polling subscriber firewall reconciler implementing both block and restore.',
    'setup/Working Backup Pending Blocker.txt':'Historical backup of the enforcement script; not the current launched module.',
    'setup/mikroktek cmd.txt':'Manual RouterOS notes for inspecting or removing IDS-owned rules.',
    'setup/Training/train_netflow_model.py':'Maintained fourteen-feature training, validation, serialization, and test metrics pipeline.',
    'setup/Training/training_schema.py':'Ordered feature names and mikrotik-netflow-v1 schema identifier.',
    'tools/build_system_documentation.py':'Offline text/PDF documentation generator; not part of website request handling.',
}
def purpose(path):
    rel=path.relative_to(ROOT).as_posix();name=path.name;suffix=path.suffix.lower()
    if rel in purpose_map:return purpose_map[rel]
    if '__pycache__' in rel:return 'Generated Python bytecode cache; recreated by Python and not authoritative source.'
    if '.ipynb_checkpoints' in rel:return 'Automatically saved historical notebook checkpoint; not an active runtime service.'
    if rel.startswith('app/Http/Controllers/'):return 'Laravel request controller; methods and behavior documented in the controller register.'
    if rel.startswith('app/Models/'):return 'Eloquent persistence model with fillable fields, casting, accessors, or relationships.'
    if rel.startswith('app/Http/Middleware/'):return 'Custom website session and role authorization middleware.'
    if rel.startswith('app/Providers/'):return 'Application service-provider lifecycle hooks.'
    if rel.startswith('database/migrations/'):return 'Versioned database schema or historical data transformation; see migration register.'
    if rel.startswith('database/seeders/'):return 'Database initialization or dedicated account seed operation.'
    if rel.startswith('database/factories/'):return 'Framework test/sample account data factory.'
    if suffix in ['.sqlite','.db']:return 'Local database artifact; existence does not prove production uses this connection.'
    if rel.startswith('config/'):return 'Laravel '+path.stem+' configuration; consumes environment settings as registered earlier.'
    if rel.startswith('routes/'):return 'HTTP or console route registration; determines exposed handler and middleware mapping.'
    if rel.startswith('bootstrap/'):return 'Laravel application or service-provider bootstrap.'
    if rel.startswith('resources/views/'):return 'Server-rendered Blade website page or partial; page responsibilities are explained in the website chapter.'
    if rel.startswith('resources/'):return 'Frontend source asset consumed by the Vite build.'
    if rel.startswith('tests/'):return 'Regression test source or shared test harness; coverage is documented in the testing chapter.'
    if 'Dataset/' in rel and suffix=='.csv':return 'Offline labeled traffic training/evaluation input; not required for ordinary inference startup.'
    if suffix=='.joblib':return 'Serialized model artifact or placeholder; inventoried without executing/deserializing. Only the configured filename is loaded at runtime.'
    if name.endswith('.metrics.json'):return 'Stored evaluation metrics and bootstrap-validation warning for a model artifact.'
    if suffix=='.docx':return 'Historical research manuscript; its textual contents are preserved in the manuscript appendix.'
    if suffix=='.ipynb':return 'Offline training or inference experiment notebook; summarized in the notebook register.'
    if name=='requirements.txt':return 'Python runtime/development dependency declaration.'
    if rel.startswith('setup/') and suffix=='.md':return 'Setup or model-development guidance.'
    if suffix=='.zip':return 'Packaged historical API or website template distribution; not automatically executed.'
    if suffix=='.exe':return 'Optional ngrok tunneling executable; no main service automatically launches it.'
    if 'PHPMailer' in rel:return 'Bundled third-party mail library or supporting metadata; no active IDS mail workflow established.'
    if suffix in ['.png','.jpg','.jpeg','.webp','.gif','.ico','.svg']:return 'Static website image, branding, illustration, favicon, or vendor icon asset.'
    if suffix in ['.ttf','.woff','.woff2','.eot','.otf']:return 'Static web-font or icon-font binary used by website styling.'
    if suffix=='.scss':return 'Template styling source or Sass mixin, variable, component, or style test; not detector logic.'
    if suffix=='.css':return 'Prebuilt website or template stylesheet.'
    if suffix=='.js':return 'Browser interaction, charting, template, or frontend tooling JavaScript.'
    if suffix in ['.html','.htm']:return 'Static website template, demonstration page, or distribution placeholder; not necessarily a Laravel route.'
    if suffix=='.map':return 'Frontend source-map metadata for debugging compiled assets.'
    if suffix=='.json':return 'Dependency, template, icon, or tooling metadata; context determined by its directory.'
    if suffix=='.php':return 'Supporting PHP library, request bootstrap, or template distribution file.'
    if suffix in ['.md','.txt']:return 'Supporting documentation, distribution notes, or license text.'
    return 'Supporting project or template asset; not a separately launched IDS component.'
groups={}
for path in files:
    rel=path.relative_to(ROOT).as_posix()
    group=rel.split('/')[0] if '/' in rel else 'Root project files'
    groups.setdefault(group,[]).append(f'{rel}. {path.stat().st_size:,} bytes. {purpose(path)}')
for group,entries in sorted(groups.items()):section('File inventory: '+group,'\n\n'.join(entries),2)
section('Grouped private, dependency, and generated directories','\n\n'.join(private_or_generated),2)
section('Documentation output files', '''
public/about_system/IDMS_Complete_System_Documentation.pdf is the final searchable, text-only PDF. public/about_system/IDMS_Complete_System_Documentation.txt is the editable plain-text companion. public/about_system/documentation_manifest.json records source hashes, source coverage counts, exclusions, build metadata, and PDF validation facts. These artifacts are generated documentation, not additional website or detection services.
''',2)

section('26. Traceability, exclusions, and documentation completion', f'''
This reference individually inventories {len(files)} project files outside the explicitly grouped private, dependency, and generated directories. It registers current Python source members, custom PHP methods, every migration's schema declarations, source-referenced environment variable names, notebook experiments, historical API archive members, and text from all five Word manuscripts. It provides a comprehensive local-source reference while keeping runtime secrets and subscriber records out of a potentially public folder.

Source-relative paths identify evidence locations. Line numbers refer to the inspected snapshot and can change after editing. The manifest stores SHA-256 fingerprints of selected first-party text sources, without copying private values. The PDF is generated from the narrative in this script plus safe local-source extraction. This document's build process reads source and development documents only; it does not contact a router, analyze live traffic, call a hosted classifier, or alter production data.

The reader should distinguish four evidence levels: behavior explicitly implemented in current source; historical claims preserved from project research; defaults that may be overridden privately; and operational outcomes requiring a live verification. A deployment is complete only when its website, schema, Python services, inference model, environment settings, and router policy agree. This reference documents the repository and known workflow; it does not replace that deployment verification.
''')

# Produce both plain text and a compact, text-only PDF with a real page-numbered contents list.
fontdir=Path('C:/Windows/Fonts')
for face,filename in [('DocRegular','arial.ttf'),('DocBold','arialbd.ttf'),('DocItalic','ariali.ttf')]:
    pdfmetrics.registerFont(TTFont(face,str(fontdir/filename)))
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyDense',fontName='DocRegular',fontSize=9.2,leading=11.5,spaceAfter=4.3,alignment=TA_JUSTIFY,splitLongWords=True))
styles.add(ParagraphStyle(name='HeadingMain',fontName='DocBold',fontSize=13,leading=15.5,spaceBefore=10,spaceAfter=5,keepWithNext=True))
styles.add(ParagraphStyle(name='HeadingSub',fontName='DocBold',fontSize=10,leading=12,spaceBefore=7,spaceAfter=3.5,keepWithNext=True))
styles.add(ParagraphStyle(name='DocTitle',fontName='DocBold',fontSize=19,leading=22,spaceAfter=7))
styles.add(ParagraphStyle(name='SmallNote',fontName='DocItalic',fontSize=8.5,leading=10.5,spaceAfter=5))


class DocumentationPDF(BaseDocTemplate):
    def __init__(self,filename):
        super().__init__(str(filename),pagesize=A4,rightMargin=32,leftMargin=32,topMargin=34,bottomMargin=32,title='IDMS Complete System Documentation',author='Project documentation',pageCompression=1)
        frame=Frame(self.leftMargin,self.bottomMargin,self.width,self.height,id='normal',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
        self.addPageTemplates(PageTemplate(id='text',frames=[frame],onPage=self.decorate))
    def decorate(self,canvas,doc):
        canvas.saveState();canvas.setFont('DocRegular',7);canvas.setFillColor(colors.HexColor('#444444'))
        canvas.drawString(32,A4[1]-20,'INTRUSION DETECTION AND MONITORING SYSTEM | SOURCE REFERENCE')
        canvas.drawString(32,18,'29 October 2026 | Text-only documentation | Private values omitted')
        canvas.drawRightString(A4[0]-32,18,f'Page {doc.page}');canvas.restoreState()
    def afterFlowable(self,flowable):
        if isinstance(flowable,Paragraph) and flowable.style.name in ['HeadingMain','HeadingSub']:
            level=0 if flowable.style.name=='HeadingMain' else 1
            label=flowable.getPlainText()
            bookmark='section_'+hashlib.sha1(label.encode()).hexdigest()
            self.canv.bookmarkPage(bookmark)
            self.canv.addOutlineEntry(label,bookmark,level=level,closed=level==0)
            if level==0:self.notify('TOCEntry',(0,label,self.page,bookmark))


plain=['INTRUSION DETECTION AND MONITORING SYSTEM\nCOMPLETE SYSTEM DOCUMENTATION\n29 October 2026\n']
story=[Paragraph('Intrusion Detection and Monitoring System',styles['DocTitle']),Paragraph('Complete system documentation: architecture, flow, Python services, website, APIs, database, model training, operations, and source inventory.',styles['BodyDense']),Paragraph('Prepared 29 October 2026. Local-source reference; deployed state requires separate verification. Text only; no private credentials.',styles['SmallNote'])]
toc=TableOfContents();toc.levelStyles=[ParagraphStyle(name='Contents',fontName='DocRegular',fontSize=8.7,leading=11,leftIndent=0,firstLineIndent=0,spaceBefore=1)]
story.append(toc);story.append(Spacer(1,8))
for level,title,paragraphs in sections:
    story.append(Paragraph(escape(title),styles['HeadingMain' if level==1 else 'HeadingSub']))
    plain.extend(['\n'+title+'\n',*paragraphs])
    for p in paragraphs:
        cleaned=''.join(c for c in p if ord(c)>=32 or c in '\n\t')
        story.append(Paragraph(escape(cleaned),styles['BodyDense']))
text='\n\n'.join(plain)
(OUT/'IDMS_Complete_System_Documentation.txt').write_text(text,encoding='utf-8')
pdfpath=OUT/'IDMS_Complete_System_Documentation.pdf'
DocumentationPDF(pdfpath).multiBuild(story)

from pypdf import PdfReader
reader=PdfReader(str(pdfpath));extracted=[p.extract_text() or '' for p in reader.pages]
source_hashes={}
for path in files:
    rel=path.relative_to(ROOT).as_posix()
    if path.suffix in ['.py','.php','.blade','.md'] and not rel.startswith(('public/','vendor/')):
        source_hashes[rel]=hashlib.sha256(path.read_bytes()).hexdigest()
manifest={
    'document_date':'2026-10-29','timezone':'Asia/Manila','scope':'Local source snapshot; no production verification',
    'project_files_individually_inventoried':len(files),'chapters':26,'subsections':len(sections)-26,
    'word_count':len(text.split()),'pdf_pages':len(reader.pages),'pdf_bytes':pdfpath.stat().st_size,
    'pages_with_no_extractable_text':[i+1 for i,t in enumerate(extracted) if not t.strip()],
    'pdf_image_objects':sum(len(p.images) for p in reader.pages),
    'private_environment_values_included':False,'serialized_models_loaded':False,
    'grouped_exclusions':private_or_generated,'source_sha256':source_hashes,
}
(OUT/'documentation_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k not in ['source_sha256','grouped_exclusions']},indent=2))
